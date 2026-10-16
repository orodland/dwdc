"""Deterministic Excel extraction (cell values read exactly; no OCR risk)."""

from __future__ import annotations

import datetime as dt
import re
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from ..config import BrokerConfig
from ..model import Cell, DocIssue, DocumentExtraction, UNCLASSIFIED_CONTENT
from .common import build_row, compile_all, is_ref_ok, parse_control_number


def cell_text(value) -> tuple[str, str]:
    """Return (raw text, value type) for an openpyxl cell value."""
    if value is None:
        return "", "empty"
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), "bool"
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0):
            return value.date().isoformat(), "date"
        return value.isoformat(), "datetime"
    if isinstance(value, dt.date):
        return value.isoformat(), "date"
    if isinstance(value, int):
        return str(value), "number"
    if isinstance(value, float):
        # repr() is the shortest round-tripping decimal form of the stored double.
        return format(Decimal(repr(value)), "f"), "number"
    return str(value).strip(), "text"


def extract_xlsx(path: Path, cfg: BrokerConfig, doc: DocumentExtraction) -> None:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        total_rows = 0
        for table in cfg.tables:
            if table.sheet not in wb.sheetnames:
                doc.issues.append(DocIssue(UNCLASSIFIED_CONTENT, f"expected sheet '{table.sheet}' not found"))
                continue
            ws = wb[table.sheet]
            end_re = compile_all(table.end_patterns)
            skip_re = compile_all(table.skip_patterns)
            totals_re = [(t, re.compile(t.regex)) for t in table.totals]
            header_map: dict[str, int] | None = None
            ctl = {"rows": 0, "quantity_sum": Decimal(0), "quantity_unparsed": 0, "footer": None}
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                texts = [cell_text(v) for v in row]
                nonempty = [t for t, _ in texts if t]
                if header_map is None:
                    stripped = [t for t, _ in texts]
                    if all(lbl in stripped for lbl in table.header):
                        header_map = {lbl: stripped.index(lbl) for lbl in table.header}
                    continue
                if not nonempty:
                    continue
                line = " ".join(nonempty)
                matched_total = False
                for rule, rx in totals_re:
                    m = rx.search(line)
                    if m:
                        ctl["footer"] = {k: v for k, v in m.groupdict().items()}
                        ctl["footer_loc"] = {"sheet": table.sheet, "row": r_idx}
                        ctl["quantity_abs"] = rule.quantity_abs
                        matched_total = True
                if any(rx.search(line) for rx in end_re):
                    break
                if matched_total:
                    continue
                if any(rx.search(line) for rx in skip_re):
                    continue
                by_label: dict[str, Cell] = {}
                for lbl, c_idx in header_map.items():
                    raw, typ = texts[c_idx] if c_idx < len(texts) else ("", "empty")
                    coord = f"{get_column_letter(c_idx + 1)}{r_idx}"
                    by_label[lbl] = Cell(raw=raw, loc={"sheet": table.sheet, "cell": coord, "type": typ})
                ref_label = table.columns.get("internal_id") or table.columns.get("broker_ref")
                if ref_label and not is_ref_ok(table, by_label[ref_label].raw):
                    doc.issues.append(
                        DocIssue(UNCLASSIFIED_CONTENT, f"row not recognised as a trade: {line[:120]}",
                                 {"sheet": table.sheet, "row": r_idx})
                    )
                    continue
                row_id = f"{table.sheet}!{r_idx}"
                er = build_row(row_id, table, by_label, {}, {"sheet": table.sheet, "row": r_idx})
                doc.rows.append(er)
                ctl["rows"] += 1
                qcell = er.cells.get("quantity")
                if qcell is not None:
                    q = parse_control_number(qcell.raw, cfg)
                    if q is None:
                        ctl["quantity_unparsed"] += 1
                    else:
                        ctl["quantity_sum"] += abs(q)
            if header_map is None:
                doc.issues.append(DocIssue(UNCLASSIFIED_CONTENT, f"header row not found on sheet '{table.sheet}'"))
            total_rows += ctl["rows"]
            doc.controls.setdefault("tables", {})[table.name] = ctl
        doc.controls["rows_extracted"] = total_rows
    finally:
        wb.close()
