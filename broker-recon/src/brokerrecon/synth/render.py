"""Render synthetic broker statements in deliberately different layouts.

BRK_A  Excel, clean single sheet, broker product codes, MONYY months, AP flag
BRK_B  Excel, title block + merged cells, two sheets, signed quantities,
       grains quoted in USD (not cents), DD.MM.YYYY text dates
BRK_C  Native PDF without grid lines, multi-page, account sections, subtotal
       lines, single-digit-year month codes, grains in eighths ("442'2")
BRK_D  Native PDF with grid lines, landscape, FX forwards, European number
       format (1.250.000,00) and DD-Mon-YYYY dates
BRK_E  CSV export with unfamiliar headers - no approved template exists yet
BRK_F  No confirmation received
"""

from __future__ import annotations

import csv
import datetime as dt
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas

from .generator import AS_OF, BrokerRow, Generator

CODES_A = {"ZC": "C", "ZW": "W", "ZS": "S", "CL": "CL", "NG": "NG", "B": "BRN", "KC": "KC", "SB": "SB", "GC": "GC", "HG": "HG"}
EXCH_A = {"ZC": "CBT", "ZW": "CBT", "ZS": "CBT", "CL": "NYM", "NG": "NYM", "B": "ICE", "KC": "ICE", "SB": "ICE", "GC": "CMX", "HG": "CMX"}
CODES_B = {"ZC": "CORN", "ZW": "WHEAT", "ZS": "SOYBEAN", "CL": "WTI", "NG": "NATGAS", "B": "BRENT", "KC": "COFFEE",
           "SB": "SUGAR11", "GC": "GOLD", "HG": "COPPER"}
GRAINS = {"ZC", "ZW", "ZS"}
MON = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]
CODE = "FGHJKMNQUVXZ"


def _d(x: Decimal) -> str:
    s = format(x.normalize(), "f")
    return "0" if s == "-0" else s


def _num(x: Decimal) -> float | int:
    return int(x) if x == x.to_integral_value() else float(_d(x))


def _ov(r: BrokerRow, fld: str, default):
    return r.overrides[fld] if fld in r.overrides else default


def _eu_group(ip: str) -> str:
    return f"{int(ip):,}".replace(",", ".")


def _eighths(price: Decimal) -> str:
    whole = int(price)
    frac = price - whole
    return f"{whole}'{int(frac * 8)}"


# ------------------------------------------------------------------------ BRK_A
def render_a(rows: list[BrokerRow], path: Path) -> int:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Open Trades"
    hdr = ["Trade Ref", "Account", "Exch", "Contract", "Mth", "Type", "P/C", "Strike", "B/S", "Lots", "Trade Price", "Trade Date", "AP"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True)
    lots = Decimal(0)
    for r in rows:
        t = r.trade
        y, m = t.contract_month.split("-")
        lots += t.quantity
        ws.append([
            r.broker_ref, t.account, EXCH_A[t.root], _ov(r, "instrument_code", CODES_A[t.root]),
            f"{MON[int(m) - 1]}{y[2:]}", "O" if t.product_type == "OPT" else "F",
            t.put_call or None, _num(t.strike) if t.strike is not None else None,
            "B" if t.side == "BUY" else "S",
            _ov(r, "quantity", _num(t.quantity)) or None, (_ov(r, "price", _num(t.price)) if "price" not in r.overrides else None),
            dt.datetime.combine(t.trade_date, dt.time()), "Y" if r.aps_flag else None,
        ])
        ws.cell(row=ws.max_row, column=12).number_format = "yyyy-mm-dd"
    ws.append([])
    ws.append(["TOTAL OPEN TRADES", len(rows), None, None, None, None, None, None, None, int(lots)])
    wb.save(path)
    return len(rows)


# ------------------------------------------------------------------------ BRK_B
def render_b(rows: list[BrokerRow], path: Path) -> dict:
    wb = openpyxl.Workbook()
    fo = wb.active
    fo.title = "Futures & Options"
    fx = wb.create_sheet("FX Forwards")
    counts = {}
    for ws, kind in ((fo, "FO"), (fx, "FX")):
        ws["A1"] = "Bravo Clearing AG (synthetic) - Statement of Open Positions"
        ws["A1"].font = Font(bold=True, size=13)
        ws.merge_cells("A1:I1")
        ws["A2"] = "Client: Synthetic Trading Co"
        ws["A3"] = f"As of: {AS_OF.strftime('%d.%m.%Y')}"
        if kind == "FO":
            ws.append([])
            ws.append(["Ticket", "Acct", "Product Code", "Contract Month", "Call/Put", "Strike Px", "Quantity", "Price", "Trade Dt"])
            sel = [r for r in rows if r.trade.product_type in ("FUT", "OPT")]
            for r in sel:
                t = r.trade
                y, m = t.contract_month.split("-")
                div = Decimal(100) if t.root in GRAINS else Decimal(1)
                q = t.quantity if t.side == "BUY" else -t.quantity
                q_out = _ov(r, "quantity", None)
                if q_out is not None and t.side == "SELL":
                    q_out = "-" + q_out
                ws.append([
                    r.broker_ref, t.account, _ov(r, "instrument_code", CODES_B[t.root]),
                    f"{MON[int(m) - 1].title()}-{y}", {"C": "Call", "P": "Put"}.get(t.put_call),
                    _num(t.strike / div) if t.strike is not None else None,
                    q_out if q_out is not None else _num(q),
                    None if "price" in r.overrides else _num(t.price / div),
                    t.trade_date.strftime("%d.%m.%Y"),
                ])
        else:
            ws.append([])
            ws.append(["Ticket", "Acct", "Currency Pair", "Buy/Sell", "Notional", "Fwd Rate", "Trade Dt", "Value Dt"])
            sel = [r for r in rows if r.trade.product_type == "FXFWD"]
            for r in sel:
                t = r.trade
                ws.append([
                    r.broker_ref, t.account, f"{t.ccy_pair[:3]}/{t.ccy_pair[3:]}", "Buy" if t.side == "BUY" else "Sell",
                    _ov(r, "quantity", _num(t.quantity)), None if "price" in r.overrides else _num(t.price),
                    t.trade_date.strftime("%d.%m.%Y"), t.value_date.strftime("%d.%m.%Y"),
                ])
        ws.append([])
        ws.append([f"Total records: {len(sel)}"])
        counts[kind] = len(sel)
    wb.save(path)
    return counts


# ------------------------------------------------------------------------ BRK_C
C_COLS = [("Ref", 40), ("Trade Date", 112), ("B/S", 175), ("Qty", 205), ("Contract", 245), ("Strike", 305), ("P/C", 360), ("Price", 400)]


def render_c(rows: list[BrokerRow], path: Path, declared_extra: int = 0) -> int:
    w, h = A4
    per_page = 52
    # Layout pass: build a list of "lines" then paginate.
    lines: list[tuple[str, object]] = []
    by_acct: dict[str, list[BrokerRow]] = {}
    for r in rows:
        by_acct.setdefault(r.trade.account, []).append(r)
    total_lots = Decimal(0)
    for acct in sorted(by_acct):
        lines.append(("section", f"Account: {acct}"))
        lines.append(("header", None))
        for r in by_acct[acct]:
            lines.append(("row", r))
            total_lots += r.trade.quantity
        lines.append(("text", f"Subtotal {acct}: {len(by_acct[acct])} trades"))
    lines.append(("text", f"Total open trades: {len(rows) + declared_extra}   Total lots: {int(total_lots):,}"))
    pages: list[list] = []
    cur: list = []
    for ln in lines:
        if len(cur) >= per_page:
            pages.append(cur)
            cur = []
        if not cur and ln[0] == "row":
            cur.append(("header", None))
        cur.append(ln)
    if cur:
        pages.append(cur)
    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    for pno, page in enumerate(pages, start=1):
        c.setFont("Helvetica-Bold", 11)
        c.drawString(40, h - 40, "Charlie Markets Ltd (synthetic) - Open Trade Statement")
        c.setFont("Helvetica", 8)
        c.drawString(40, h - 54, f"Statement date {AS_OF.strftime('%d/%m/%Y')}   Client: Synthetic Trading Co")
        y = h - 80
        for kind, payload in page:
            if kind == "header":
                c.setFont("Helvetica-Bold", 8)
                for label, x in C_COLS:
                    c.drawString(x, y, label)
                c.setFont("Helvetica", 8)
            elif kind == "section":
                c.setFont("Helvetica-Bold", 8)
                c.drawString(40, y, payload)
                c.setFont("Helvetica", 8)
            elif kind == "text":
                c.drawString(40, y, payload)
            else:
                r: BrokerRow = payload
                t = r.trade
                y_, m = t.contract_month.split("-")
                price = "" if "price" in r.overrides else (_eighths(t.price) if t.root in GRAINS else _d(t.price))
                vals = [r.broker_ref, t.trade_date.strftime("%d/%m/%Y"), "B" if t.side == "BUY" else "S",
                        _ov(r, "quantity", _d(t.quantity)),
                        f"{_ov(r, 'instrument_code', t.root)} {CODE[int(m) - 1]}{y_[3]}",
                        _d(t.strike) if t.strike is not None else "", t.put_call or "", price]
                for (label, x), v in zip(C_COLS, vals):
                    if v:
                        c.drawString(x, y, v)
            y -= 13
        c.setFont("Helvetica-Oblique", 7)
        c.drawString(40, 30, "Synthetic statement for prototype testing only. Not a real document.")
        c.drawRightString(w - 40, 30, f"Page {pno} of {len(pages)}")
        c.showPage()
    c.save()
    return len(rows)


# ------------------------------------------------------------------------ BRK_D
D_COLS = [("Deal No", 40), ("Account", 125), ("Pair", 195), ("Buy/Sell", 255), ("Base Amount", 320), ("Forward Rate", 440),
          ("Trade Date", 540), ("Value Date", 640)]


def render_d(rows: list[BrokerRow], path: Path) -> int:
    w, h = landscape(A4)
    per_page = 30  # leaves room for the closing control-total line above the page footer
    pages = [rows[i:i + per_page] for i in range(0, len(rows), per_page)] or [[]]
    c = canvas.Canvas(str(path), pagesize=landscape(A4), invariant=1)
    for pno, chunk in enumerate(pages, start=1):
        c.setFont("Helvetica-Bold", 12)
        c.drawString(40, h - 40, "Delta FX & Commodities SA (synthetic)")
        c.setFont("Helvetica", 8)
        c.drawString(40, h - 54, f"Confirmation of outstanding forward deals at {AS_OF.strftime('%d-%b-%Y')}")
        y = h - 85
        c.setLineWidth(0.4)
        c.rect(36, y - 4, w - 72, 14, stroke=1, fill=0)
        c.setFont("Helvetica-Bold", 8)
        for label, x in D_COLS:
            c.drawString(x, y, label)
        c.setFont("Helvetica", 8)
        y -= 14
        for r in chunk:
            t = r.trade
            dp = 3 if t.ccy_pair.endswith("JPY") else 5
            vals = [r.broker_ref, t.account, f"{t.ccy_pair[:3]}/{t.ccy_pair[3:]}", "Buy" if t.side == "BUY" else "Sell",
                    _ov(r, "quantity", _eu_group(str(int(t.quantity))) + ",00"),
                    "" if "price" in r.overrides else f"{t.price:.{dp}f}".replace(".", ","),
                    t.trade_date.strftime("%d-%b-%Y"), t.value_date.strftime("%d-%b-%Y")]
            c.rect(36, y - 4, w - 72, 14, stroke=1, fill=0)
            for (label, x), v in zip(D_COLS, vals):
                if v:
                    c.drawString(x, y, v)
            y -= 14
        if pno == len(pages):
            c.drawString(40, y - 10, f"Number of deals: {len(rows)}")
        c.setFont("Helvetica-Oblique", 7)
        c.drawString(40, 25, "Synthetic statement for prototype testing only. Not a real document.")
        c.drawRightString(w - 40, 25, f"Page {pno} of {len(pages)}")
        c.showPage()
    c.save()
    return len(rows)


# ------------------------------------------------------------------------ BRK_E
def render_e(rows: list[BrokerRow], path: Path) -> int:
    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["ExecID", "Client Account", "Symbol", "Expiry", "Side", "Size", "Px", "Exec Date"])
        for r in rows:
            t = r.trade
            w.writerow([r.broker_ref, t.account, _ov(r, "instrument_code", t.root), t.contract_month, t.side,
                        _ov(r, "quantity", _d(t.quantity)), "" if "price" in r.overrides else _d(t.price), t.trade_date.isoformat()])
        w.writerow([f"# records: {len(rows)}"])
    return len(rows)


# ------------------------------------------------------------------- internal
INTERNAL_HDR = ["InternalTradeID", "BrokerID", "Account", "ProductType", "Exchange", "Root", "ContractMonth", "PutCall",
                "Strike", "CcyPair", "ValueDate", "Side", "Quantity", "Price", "TradeDate", "BrokerRef", "APSGroupID"]


def render_internal(g: Generator, path: Path) -> int:
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("OpenTrades")
    ws.append(INTERNAL_HDR)
    exch = {"ZC": "XCBT", "ZW": "XCBT", "ZS": "XCBT", "CL": "XNYM", "NG": "XNYM", "B": "IFEU", "KC": "IFUS", "SB": "IFUS", "GC": "XCEC", "HG": "XCEC"}
    trades = sorted(g.trades, key=lambda t: t.internal_id)
    for t in trades:
        ws.append([
            t.internal_id, t.broker_id, t.account, t.product_type, exch.get(t.root) if t.root else None, t.root,
            t.contract_month, t.put_call, _num(t.strike) if t.strike is not None else None, t.ccy_pair,
            dt.datetime.combine(t.value_date, dt.time()) if t.value_date else None, t.side, _num(t.quantity),
            _num(t.price), dt.datetime.combine(t.trade_date, dt.time()),
            t.broker_ref if t.internal_has_ref else None, t.aps_group_id,
        ])
    ws.append([])
    ws.append([f"Total records: {len(trades)}"])
    wb.save(path)
    return len(trades)


def render_all(g: Generator, evidence: Path) -> tuple[dict, dict]:
    b = evidence / "brokers"
    docs = {}
    rows = g.rows
    docs["brokers/alpha_open_trades_20261231.xlsx"] = {"broker_id": "BRK_A", "rows": render_a(rows["BRK_A"], b / "alpha_open_trades_20261231.xlsx"), "doc_codes": []}
    counts = render_b(rows["BRK_B"], b / "bravo_statement_31122026.xlsx")
    docs["brokers/bravo_statement_31122026.xlsx"] = {"broker_id": "BRK_B", "rows": sum(counts.values()), "doc_codes": []}
    # BRK_C: two statements, one per account group; the second declares one trade too many.
    c_rows = rows["BRK_C"]
    accts = sorted({r.trade.account for r in c_rows})
    s1 = [r for r in c_rows if r.trade.account == accts[0]]
    s2 = [r for r in c_rows if r.trade.account != accts[0]]
    docs["brokers/charlie_stmt_2026-12-31_part1.pdf"] = {"broker_id": "BRK_C", "rows": render_c(s1, b / "charlie_stmt_2026-12-31_part1.pdf"), "doc_codes": []}
    docs["brokers/charlie_stmt_2026-12-31_part2.pdf"] = {"broker_id": "BRK_C", "rows": render_c(s2, b / "charlie_stmt_2026-12-31_part2.pdf", declared_extra=1),
                                                         "doc_codes": ["COMPLETENESS_FAIL"]}
    docs["brokers/delta_fx_confirmation.pdf"] = {"broker_id": "BRK_D", "rows": render_d(rows["BRK_D"], b / "delta_fx_confirmation.pdf"), "doc_codes": []}
    docs["brokers/echo_positions_export.csv"] = {"broker_id": "BRK_E", "rows": render_e(rows["BRK_E"], b / "echo_positions_export.csv"),
                                                 "doc_codes_before_onboarding": ["FORMAT_UNRECOGNISED"], "doc_codes": []}
    render_internal(g, evidence / "internal" / "internal_open_trades.xlsx")
    intake = {
        "note": "Confirmation intake register maintained by the audit team (synthetic). Only files listed here are processed.",
        "internal": {"file": "internal/internal_open_trades.xlsx", "provided_by": "client finance (synthetic)",
                     "description": "Open derivative trades extract from trading system as at year end"},
        "confirmations": [
            {"file": f, "broker_id": d["broker_id"], "received_from": f"{d['broker_id']} confirmations desk",
             "channel": "auditor-controlled confirmation platform", "received_on": "2027-01-08"}
            for f, d in docs.items()
        ],
        "requested_not_received": [{"broker_id": "BRK_F", "requested_on": "2027-01-04"}],
    }
    return docs, intake
